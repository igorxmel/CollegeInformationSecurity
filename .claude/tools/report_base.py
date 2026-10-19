"""Оформление отчётов курса в .docx.

Параметры страницы и абзаца повторяют отчёты по бизнес-аналитике: тот же свод правил,
Times New Roman 14, одинарный интервал, отступ первой строки 1,5 см, без жирного,
курсива и подчёркивания, без разрывов страницы между разделами.

Отличие одно: здесь есть титульный лист.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
SIZE = Pt(14)
SMALL = Pt(11)
INDENT = Cm(1.5)
TEXT_WIDTH = 16.5
DASH = "–"

LINES_PER_PAGE = 45

TITLE = [
    "Министерство общего и профессионального образования",
    "Ростовской области",
    "Государственное бюджетное профессиональное образовательное",
    "учреждение Ростовской области",
    "«Ростовский-на-Дону колледж связи и информатики»",
]
SIGN_LEFT = ["Выполнил: студент 4 курса", "группы ИС-42", "Холкин Егор"]
SIGN_RIGHT = ["Проверил: преподаватель", "Бойнар И. Н."]
CITY = "Ростов-на-Дону"
YEAR = "2026"


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.font.bold = False
    normal.font.italic = False
    normal.font.underline = False
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    pf = normal.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc


def para(doc, text="", *, indent=INDENT, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = indent
    p.paragraph_format.alignment = align
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = SIZE
    return p


def blank(doc, count=1):
    for _ in range(count):
        para(doc)


def section_heading(doc, number, title, *, first=False):
    """Раздел без разрыва страницы: одна пустая строка до, две после."""
    if not first:
        blank(doc, 1)
    para(doc, f"{number} {title}")
    blank(doc, 2)


def sub_heading(doc, number, title):
    blank(doc, 1)
    para(doc, f"{number} {title}")
    blank(doc, 2)


def bullets(doc, items):
    for item in items:
        para(doc, f"{DASH} {item}")


def table(doc, number, caption, header, rows, widths, *, size=SIZE):
    if abs(sum(widths) - TEXT_WIDTH) > 0.05:
        raise SystemExit(f"ширина таблицы {sum(widths):.2f} см вместо {TEXT_WIDTH}")

    blank(doc, 1)
    para(doc, f"Таблица {number} {DASH} {caption}")

    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False

    def fill(cells, values, center):
        for cell, value, width in zip(cells, values, widths):
            cell.width = Cm(width)
            cell.text = ""
            for i, line in enumerate(str(value).split("\n")):
                p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if center
                               else WD_ALIGN_PARAGRAPH.LEFT)
                run = p.add_run(line)
                run.font.name = FONT
                run.font.size = size

    fill(t.rows[0].cells, header, True)
    for row in rows:
        fill(t.add_row().cells, row, False)

    blank(doc, 1)


def title_page(doc, number, topic):
    """Титульный лист на всю первую страницу, дальше разрыв.

    Пустые строки подобраны так, чтобы блоки разошлись по высоте листа: шапка
    сверху, тема по центру, подписи ниже, город и год внизу. Разрыв страницы
    здесь единственный в документе, между разделами их нет.
    """
    lines = 0
    for line in TITLE:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    lines += len(TITLE)

    blank(doc, 8)
    lines += 8

    para(doc, "ОТЧЁТ", indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, f"по практической работе № {number}", indent=Cm(0),
         align=WD_ALIGN_PARAGRAPH.CENTER)
    blank(doc)
    topic_lines = max(1, -(-len(topic) // 67))
    para(doc, topic, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    lines += 3 + topic_lines

    blank(doc, 10)
    lines += 10

    for line in SIGN_LEFT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    blank(doc)
    for line in SIGN_RIGHT:
        para(doc, line, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    lines += len(SIGN_LEFT) + 1 + len(SIGN_RIGHT)

    blank(doc, 8)
    lines += 8

    para(doc, CITY, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, YEAR, indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    lines += 2

    if lines > LINES_PER_PAGE:
        raise SystemExit(f"титульный лист занял {lines} строк при лимите {LINES_PER_PAGE}")

    para(doc).add_run().add_break(WD_BREAK.PAGE)


def save(doc, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(path)
    except PermissionError:
        raise SystemExit(f"{path.name} открыт в Word, закрой его и повтори")
    print(f"saved: {path.name}")
